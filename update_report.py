import shutil
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

URL = "https://khoachuangur.github.io/Khoa_PhamDang_526000180/"
SRC = "CSDL_Nhom1_DeTai15_.docx"
OUT = "CSDL_Nhom1_DeTai15_HoanChinh.docx"

shutil.copy(SRC, SRC + ".bak")
d = docx.Document(SRC)


def add_link(par, url):
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = url; r.append(t)
    h.append(r)
    par._p.append(h)


# Mục 8.4: thêm đoạn mới ngay sau đoạn nội dung
idx = next(i for i, p in enumerate(d.paragraphs) if p.text.startswith("8.4."))
body = d.paragraphs[idx + 1]
if URL not in body.text:
    new_p = OxmlElement("w:p")
    body._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    np_ = Paragraph(new_p, body._parent)
    np_.style = body.style
    np_.add_run("- Trang ePortfolio online (GitHub Pages): ")
    add_link(np_, URL)

# Phụ lục A: dòng "Nhật ký ePortfolio"
done = False
for t in d.tables:
    for row in t.rows:
        if row.cells[0].text.strip() == "Nhật ký ePortfolio":
            cell = row.cells[-1]
            p = cell.paragraphs[0]
            for r in p.runs:
                r.text = ""
            add_link(p, URL)
            done = True
assert done, "Không tìm thấy dòng Nhật ký ePortfolio"

d.save(OUT)

# Kiểm tra
c = docx.Document(OUT)
print("8.4 OK:", any(URL in "".join(x.text for x in p._p.iter(qn("w:t"))) for p in c.paragraphs))
for t in c.tables:
    for row in t.rows:
        if row.cells[0].text.strip() == "Nhật ký ePortfolio":
            print("PL A:", [x.text for x in row.cells], "|", "".join(x.text or "" for x in row._tr.iter(qn("w:t"))))
